"""Realized ownership, conserved corporate flows and frozen-result attribution."""
from collections import defaultdict
from decimal import Decimal, localcontext
import json
from pathlib import Path
import subprocess
import sys

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
import stage1_attribution as m


@pytest.fixture(scope='module')
def holdings():
    return m.read(m.RESULT/'annual_holdings_attribution.csv')


def group(holdings, p, year):
    return [r for r in holdings if r['portfolio']==p and r['end']==m.DATES[year+1]]


@pytest.mark.parametrize('year', range(2014, 2026))
@pytest.mark.parametrize('portfolio', m.PORTFOLIOS)
def test_exact_decimal_contribution_sum_at_original_precision(holdings, year, portfolio):
    annual = m.read(m.RESULT/'annual_returns_pct.csv')[year-2014]
    with localcontext() as ctx:
        ctx.prec = 50
        rows = group(holdings, portfolio, year)
        assert sum(Decimal(r['contribution_pp']) for r in rows) == Decimal(annual[portfolio])
        residue = [r for r in rows if r['row_type']=='NUMERICAL_RESIDUAL']
        assert len(residue)==1 and abs(Decimal(residue[0]['contribution_pp'])) < Decimal('1e-9')
        assert not residue[0]['exposure_return_pct'] and not residue[0]['start_ticker']


@pytest.mark.parametrize('portfolio', m.PORTFOLIOS)
def test_every_start_holding_uses_actual_opening_weight_and_pre_review_endpoint(holdings, portfolio):
    for year in range(2014, 2026):
        actual = {r['start_ticker']:r for r in group(holdings, portfolio, year) if r['row_type']=='HOLDING'}
        if portfolio=='R03 B2' and year<2016:
            expected = {r['ticker']:r for r in m.read(m.OUT/'established_segment_positions.csv') if r['year']==str(year)}
        elif portfolio=='BH padrão' or year==2014:
            filename = 'bh_june_positions.csv' if portfolio=='BH padrão' else 'b00s_initial_positions.csv'
            expected = {r['ticker']:r for r in m.read(m.OUT/filename) if r['date']==m.DATES[year]}
        else:
            expected = {r['ticker']:r for r in m.read(m.OUT/f'{m.STEMS[portfolio]}_positions.csv')
                        if r['date']==m.DATES[year] and r['phase']=='AFTER_REVIEW'}
        assert actual.keys()==expected.keys()
        for t, r in actual.items():
            assert float(r['start_weight_pct']) == pytest.approx(100*float(expected[t]['weight']), abs=1e-13)
            assert float(r['contribution_pp']) == pytest.approx(float(r['start_weight_pct'])*float(r['exposure_return_pct'])/100, abs=1e-12)
        assert sum(float(r['start_weight_pct']) for r in actual.values())==pytest.approx(100, abs=1e-10)
        assert sum(float(r['end_weight_before_review_pct']) for r in actual.values())==pytest.approx(100, abs=1e-10)


def test_new_closing_june_purchases_excluded_and_closing_exits_retained(holdings):
    count = 0
    for p, stem in m.STEMS.items():
        reviews = m.read(m.OUT/f'{stem}_reviews.csv')
        positions = m.read(m.OUT/f'{stem}_positions.csv')
        for year in range(2016 if p=='R03 B2' else 2015, 2025):
            closing = [r for r in reviews if r['date']==m.DATES[year+1]]
            opening = {r['ticker'] for r in positions if r['date']==m.DATES[year] and r['phase']=='AFTER_REVIEW'}
            actual = {r['start_ticker'] for r in group(holdings, p, year) if r['row_type']=='HOLDING'}
            next_cycle = {r['start_ticker'] for r in group(holdings, p, year+1) if r['row_type']=='HOLDING'}
            for r in closing:
                if float(r['before'])==0 and r['ticker'] not in opening:
                    assert r['ticker'] not in actual and r['ticker'] in next_cycle
                    count += 1
                if float(r['after'])==0 and r['ticker'] in opening:
                    assert r['ticker'] in actual
                    assert r['ticker'] not in next_cycle
    assert count > 10


def test_initial_attributions_reused_verbatim_without_replaying_events(holdings):
    for p in m.PORTFOLIOS:
        if p=='R03 B2':
            ready = [r for r in m.read(m.OUT/'established_segment_positions.csv') if r['year']=='2014']
            key='ticker'
        elif p=='BH padrão':
            ready = [r for r in m.read(m.OUT/'bh_issuer_annual_pct.csv') if r['year']=='2014']
            key='original_issuer_ticker'
        else:
            ready = m.read(m.OUT/'b00s_initial_contributions.csv'); key='ticker'
        actual = {r['start_ticker']: r for r in group(holdings, p, 2014) if r['row_type']=='HOLDING'}
        for r in ready:
            assert actual[r[key]]['contribution_pp']==r['contribution_pp']
            assert actual[r[key]]['exposure_return_pct']==r['return_pct']


def test_right_detaches_and_reinvests_once_in_original_exposure():
    day='2026-05-20'
    events=[dict(id='r', ticker='P', kind='RIGHT', ex_date=day, successor='R', ratio=.5),
            dict(id='s', ticker='R', kind='RIGHT_REINVEST', ex_date=day, successor='P')]
    after, transfers = m.attribute_day({'origin':{'P':2}}, events, {('P',day):10, ('R',day):2}, day)
    assert after=={'origin':{'P':2.2}} and not transfers
    with pytest.raises(ValueError, match='Duplicate'):
        m.attribute_day({'origin':{'P':2}}, events+events, {('P',day):10, ('R',day):2}, day)


def test_spinoff_separates_pre_existing_child_without_double_counting():
    day='2021-10-04'
    events=[dict(id='split', ticker='P', kind='SPINOFF', ex_date=day, successor='C', ratio=.5)]
    after, _ = m.attribute_day({'P':{'P':2}, 'C':{'C':3}}, events, {}, day)
    assert after=={'P':{'P':2, 'C':1}, 'C':{'C':3}}
    assert m.aggregate(after)=={'P':2, 'C':4}


def test_conversion_existing_successor_and_old_entitlement_stay_separate():
    day='2024-11-01'
    events=[dict(id='c', ticker='OLD', kind='CONVERSION', ex_date=day, legs=[['NEW',.6]], amount=1)]
    after, _ = m.attribute_day({'OLD':{'OLD':2}, 'NEW':{'NEW':1}}, events, {('NEW',day):5}, day)
    assert after['OLD']==pytest.approx({'NEW':1.6})
    assert after['NEW']=={'NEW':1}


def test_redemption_principal_and_later_return_stay_with_original_exposure():
    day='2023-09-13'
    events=[dict(id='redeem', ticker='OLD', kind='REDEMPTION', ex_date=day, amount=30, source='frozen')]
    after, transfers = m.attribute_day({'OLD':{'OLD':1}, 'A':{'A':2}, 'B':{'B':1}}, events,
                                       {('A',day):10, ('B',day):20}, day)
    assert after=={'OLD':{'A':1.5, 'B':.75}, 'A':{'A':2}, 'B':{'B':1}}
    assert sum(r['transferred_index_component'] for r in transfers)==30
    # If A subsequently doubles and B is flat, OLD earns 15, A earns 20,
    # B earns zero. Neither OLD's principal nor its 15 is credited to A.
    finals={o:sum(q*{'A':20,'B':20}[t] for t,q in units.items()) for o,units in after.items()}
    assert finals=={'OLD':45, 'A':40, 'B':20}
    assert sum(finals.values())-70==35


def test_actual_redemption_transfers_are_conserved_and_endpoint_sources_reconcile(holdings):
    transfers = m.read(m.RESULT/'attribution_redemption_transfers.csv')
    groups=defaultdict(list)
    for r in transfers:
        groups[r['portfolio'], r['event_id'], r['origin_ticker']].append(r)
    assert {r['redeemed_ticker'] for r in transfers}=={'ENBR3','NEOE3'}
    for rs in groups.values():
        assert sum(float(r['recipient_weight']) for r in rs)==pytest.approx(1, abs=1e-12)
        assert sum(float(r['transferred_index_component']) for r in rs)==pytest.approx(float(rs[0]['redeemed_index_component']), abs=1e-13)
    endpoints=m.read(m.RESULT/'attribution_end_lineage.csv')
    positions={p:m.read(m.OUT/f'{stem}_positions.csv') for p,stem in m.STEMS.items()}
    for p, stem in m.STEMS.items():
        for year in range(2016 if p=='R03 B2' else 2015,2026):
            rs=[r for r in endpoints if r['portfolio']==p and r['end']==m.DATES[year+1]]
            physical=defaultdict(float)
            for r in rs:
                physical[r['end_ticker']]+=float(r['index_units'])
            expected={r['ticker']:float(r['units']) for r in positions[p] if r['date']==m.DATES[year+1] and r['phase']=='PERIOD_END'}
            assert physical==pytest.approx(expected, abs=1e-14)
            source_values=defaultdict(float)
            for r in rs:
                source_values[r['origin_ticker']]+=float(r['index_component'])
            for r in group(holdings,p,year):
                if r['row_type']=='HOLDING':
                    assert source_values[r['start_ticker']]==pytest.approx(float(r['end_index_component']), abs=1e-14)


def test_xp_is_child_in_creation_cycle_then_an_opening_holding(holdings):
    for p in m.PORTFOLIOS:
        for year in (2021,2022):
            rs={r['start_ticker']:r for r in group(holdings,p,year) if r['row_type']=='HOLDING'}
            if 'ITUB4' not in rs:
                continue
            if year==2021:
                assert 'XPBR31' not in rs and 'XPBR31' in rs['ITUB4']['related_securities']
            else:
                assert 'XPBR31' in rs and 'XPBR31' not in rs['ITUB4']['related_securities']


def test_original_backtests_aggregates_ibov_and_books_byte_preserved():
    assert m.verify_frozen()==271


def test_attribution_does_not_invoke_selection_renewal_or_backtests(monkeypatch):
    import stage1_resume
    import stage1_buyhold
    import stage1_continuity
    import stage1_segments

    def forbidden(*args, **kwargs):
        pytest.fail('Attribution must not invoke a selection or backtest')

    for module, name in [(stage1_resume,'resume'), (stage1_resume,'selection'),
                         (stage1_resume,'renew_b2'), (stage1_buyhold,'trajectory'),
                         (stage1_buyhold,'run'), (stage1_continuity,'renew_b2'),
                         (stage1_segments,'gross_factor')]:
        monkeypatch.setattr(module, name, forbidden)
    rows, summaries, *_ = m.build()
    assert len(summaries)==48 and sum(r['row_type']=='HOLDING' for r in rows)==729


def test_excel_contains_all_holdings_totals_filters_and_cached_results(holdings):
    from stage1_attribution_excel import SHEETS
    path=m.RESULT/'attribution_2014_2026.xlsx'
    book=openpyxl.load_workbook(path, data_only=True)
    formulas=openpyxl.load_workbook(path, data_only=False)
    assert 'Resumo 12 anos' in book and 'Conciliacao 48 totais' in book
    published=m.read(m.RESULT/'annual_returns_pct.csv')
    for p,sheet in SHEETS.items():
        ws=book[sheet]
        assert ws.freeze_panes=='D6' and len(ws.tables)==1
        actual=list(ws.iter_rows(min_row=6, values_only=True))
        total_rows=[r for r in actual if r[11]=='TOTAL']
        assert len(total_rows)==12
        for year,r in zip(range(2014,2026),total_rows):
            assert r[6]==pytest.approx(float(published[year-2014][p]), abs=1e-13)
            assert r[10]==pytest.approx(float(published[year-2014]['IBOV']), abs=1e-13)
            assert r[14]==0
        holding_rows=[r for r in actual if r[11]=='HOLDING']
        expected=[r for r in holdings if r['portfolio']==p and r['row_type']=='HOLDING']
        assert len(holding_rows)==len(expected)
        for a,b in zip(holding_rows,expected):
            assert a[2]==b['start_ticker'] and a[12]==b['contribution_pp']
            assert a[4]==pytest.approx(float(b['start_weight_pct']), abs=1e-13)
        formula_rows=[r for r in formulas[sheet].iter_rows(min_row=6) if r[11].value=='TOTAL']
        assert all(r[6].data_type=='f' and r[14].data_type=='f' for r in formula_rows)


def test_attribution_rebuild_is_offline_byte_identical_and_does_not_run_backtests(tmp_path):
    subprocess.run([sys.executable, str(ROOT/'scripts/stage1_attribution.py'), '--output-dir', str(tmp_path)],
                   cwd=ROOT, check=True, capture_output=True)
    paths=list(tmp_path.iterdir())
    assert len(paths)==7
    for p in paths:
        assert p.read_bytes()==(m.RESULT/p.name).read_bytes(), p.name
    m.verify_frozen()
